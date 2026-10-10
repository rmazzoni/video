"""Dubbing tab for a Sleepy episode.

The commands match Main's Dubbing tab. Speech is the Qwen clone, in the
Italian or English voice selected here, written to output/audio/scene_NNN.mp3.
Speed is applied to that file after the clone speaks. Edge TTS is not used.
"""

from __future__ import annotations

import math
import os
import re
import subprocess
import tempfile
from pathlib import Path

import yaml
from PyQt6.QtCore import QEvent, QObject, Qt, QThread, QTimer, QUrl, pyqtSignal
from PyQt6.QtGui import (
    QColor,
    QKeySequence,
    QShortcut,
    QSyntaxHighlighter,
    QTextCharFormat,
    QTextCursor,
)
from PyQt6.QtMultimedia import QAudioOutput, QMediaPlayer
from PyQt6.QtWidgets import (
    QCheckBox,
    QComboBox,
    QDoubleSpinBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSlider,
    QVBoxLayout,
    QWidget,
)

from narration.qwen_voice import VoiceCancelled, VoiceRun, synthesize_scenes, validate_request
from narration.tts_engine import apply_tts_fixups
from sleepy.chapter import LANGUAGES, normalize_language
from video.ken_burns_generator import OPTIMAL_SHOT_SECONDS, required_shot_count

_GPU_OWNER = "sleepy-voice"

_NOT_DUBBED = ("#0e2a3d", "#5BB4D8")
_TO_REDUB = ("#3d3000", "#F8B23D")
_DUBBED = ("#0d2e14", "#4CAF50")
_BTN_COLORS = {
    "none": ("#555555", "#444444"),
    "all_stale": ("#c0392b", "#a93226"),
    "stale": ("#e67e22", "#ca6f1e"),
    "complete": ("#27ae60", "#1e8449"),
}
_BTN_SS = (
    "QPushButton {{ background-color:{bg}; color:white; font-weight:bold; "
    "font-size:13px; border-radius:5px; padding:4px 14px; border:none; }}"
    "QPushButton:hover {{ background-color:{hv}; }}"
    "QPushButton:pressed {{ background-color:{hv}; }}"
    "QPushButton:disabled {{ background-color:#555; color:#999; }}"
)


class _SpellHighlighter(QSyntaxHighlighter):
    """Underlines misspelled words. Same dictionary file as Main's Dubbing tab."""

    _ITALIAN_ELISIONS = {
        "all", "bell", "c", "coll", "d", "dall", "dell", "gl", "l",
        "m", "n", "nell", "quest", "quell", "s", "senz", "sott", "sull", "t",
        "un",
    }
    _CUSTOM_WORDS_PATH = os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
        "config", "spell_custom_words.txt",
    )
    _custom_words: set[str] = set()

    @classmethod
    def _load_custom_words(cls) -> set[str]:
        words: set[str] = set()
        try:
            path = Path(cls._CUSTOM_WORDS_PATH)
            lines = path.read_text(encoding="utf-8-sig").splitlines()
            cleaned: list[str] = []
            for line in lines:
                word = line.split("#", 1)[0].strip()
                normalized = word.casefold()
                if word and normalized in words:
                    continue
                cleaned.append(line)
                if word:
                    words.add(normalized)
            if cleaned != lines:
                path.write_text("\n".join(cleaned).rstrip() + "\n", encoding="utf-8")
        except FileNotFoundError:
            pass
        except Exception:
            pass
        return words

    @classmethod
    def add_custom_word(cls, word: str) -> bool:
        normalized = word.strip().casefold()
        cls._custom_words = cls._load_custom_words()
        if not normalized or normalized in cls._custom_words:
            return False
        path = Path(cls._CUSTOM_WORDS_PATH)
        path.parent.mkdir(parents=True, exist_ok=True)
        needs_newline = path.exists() and path.stat().st_size > 0
        with path.open("a", encoding="utf-8") as handle:
            if needs_newline:
                with path.open("rb") as source:
                    source.seek(-1, os.SEEK_END)
                    if source.read(1) not in (b"\n", b"\r"):
                        handle.write("\n")
            handle.write(word.strip() + "\n")
        cls._custom_words.add(normalized)
        return True

    def __init__(self, parent, language: str = "it"):
        super().__init__(parent)
        self._fmt = QTextCharFormat()
        self._fmt.setUnderlineStyle(QTextCharFormat.UnderlineStyle.SpellCheckUnderline)
        self._fmt.setUnderlineColor(QColor("#FF5555"))
        self._checker = None
        if not _SpellHighlighter._custom_words:
            _SpellHighlighter._custom_words = _SpellHighlighter._load_custom_words()
        self.set_language(language)

    def set_language(self, language: str) -> None:
        try:
            from spellchecker import SpellChecker
            self._checker = SpellChecker(language=language)
            if _SpellHighlighter._custom_words:
                self._checker.word_frequency.load_words(_SpellHighlighter._custom_words)
        except Exception:
            self._checker = None
        self.rehighlight()

    def highlightBlock(self, text: str) -> None:
        if not self._checker or not text.strip():
            return
        for match in re.finditer(r"[A-Za-zÀ-ÖØ-öø-ÿ]+(?:['’][A-Za-zÀ-ÖØ-öø-ÿ]+)?", text):
            word = match.group()
            parts = re.split(r"['’]", word, maxsplit=1)
            checked = word
            offset = 0
            if len(parts) == 2 and parts[0].casefold() in self._ITALIAN_ELISIONS:
                checked = parts[1]
                offset = len(parts[0]) + 1
            if checked.casefold() not in self._custom_words and self._checker.unknown([checked]):
                self.setFormat(match.start() + offset, len(checked), self._fmt)


def rate_percent(speed: float) -> int:
    return int(round((float(speed) - 1.0) * 100))


def speed_badge_text(percent, has_audio: bool) -> str:
    if not has_audio or percent is None:
        return "—"
    return f"{1 + int(percent) / 100:.2f}x"


def apply_playback_speed(path: str, speed: float) -> None:
    """Stretch a scene mp3. 1.0 leaves the clone's file untouched."""
    factor = float(speed)
    if abs(factor - 1.0) < 0.005:
        return
    factor = min(2.0, max(0.5, factor))
    folder = os.path.dirname(path) or "."
    handle, temporary = tempfile.mkstemp(suffix=".mp3", dir=folder)
    os.close(handle)
    try:
        result = subprocess.run(
            [
                "ffmpeg", "-y", "-i", path,
                "-filter:a", f"atempo={factor:.4f}",
                "-c:a", "libmp3lame", "-b:a", "192k",
                temporary,
            ],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0 or not os.path.isfile(temporary) or os.path.getsize(temporary) == 0:
            tail = (result.stderr or "")[-800:]
            raise RuntimeError(f"Could not apply speed {factor:.2f}x.\n{tail}")
        os.replace(temporary, path)
        temporary = ""
    finally:
        if temporary and os.path.exists(temporary):
            try:
                os.remove(temporary)
            except OSError:
                pass


def _audio_seconds(path: str) -> float:
    if not path or not os.path.isfile(path):
        return 0.0
    try:
        from mutagen.mp3 import MP3
        return float(MP3(path).info.length)
    except Exception:
        pass
    result = subprocess.run(
        [
            "ffprobe", "-v", "error", "-show_entries", "format=duration",
            "-of", "default=noprint_wrappers=1:nokey=1", path,
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode == 0 and result.stdout.strip():
        return float(result.stdout.strip())
    return 0.0


def _clock(seconds: float) -> str:
    total = max(0, int(float(seconds) + 0.5))
    minutes, secs = divmod(total, 60)
    return f"{minutes:02d}:{secs:02d}"


def _yaml_entry(existing: dict, scene_id: int) -> dict:
    if not isinstance(existing, dict):
        return {}
    entry = existing.get(scene_id, existing.get(str(scene_id), {}))
    return entry if isinstance(entry, dict) else {}


class _DubWorker(QObject):
    progress = pyqtSignal(str)
    scene_ready = pyqtSignal(int)
    finished = pyqtSignal(bool, str)

    def __init__(self, kwargs: dict, voice_run: VoiceRun, speed: float):
        super().__init__()
        self._kwargs = kwargs
        self._voice_run = voice_run
        self._speed = float(speed)

    def run(self) -> None:
        def _on_progress(message: str) -> None:
            text = str(message or "")
            self.progress.emit(text)
            if not text.startswith("Spoke scene "):
                return
            parts = text.split()
            if len(parts) < 3:
                return
            try:
                scene_id = int(parts[2])
            except ValueError:
                return
            audio = os.path.join(str(self._kwargs["output_dir"]), f"scene_{scene_id:03d}.mp3")
            try:
                apply_playback_speed(audio, self._speed)
                _store_scene_timing(str(self._kwargs["timings_path"]), scene_id, _audio_seconds(audio))
            except Exception as exc:
                self.progress.emit(f"Scene {scene_id} speed change failed: {exc}")
            self.scene_ready.emit(scene_id)

        try:
            synthesize_scenes(
                on_progress=_on_progress,
                voice_run=self._voice_run,
                **self._kwargs,
            )
            self.finished.emit(True, "Dubbing finished.")
        except VoiceCancelled as exc:
            self.finished.emit(False, str(exc))
        except Exception as exc:
            self.finished.emit(False, str(exc))


def _store_scene_timing(path: str, scene_id: int, seconds: float) -> None:
    data = {}
    if os.path.isfile(path):
        loaded = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
        if isinstance(loaded, dict):
            data = loaded
    data[int(scene_id)] = round(float(seconds), 3)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    Path(path).write_text(yaml.safe_dump(data, sort_keys=True), encoding="utf-8")


class SleepyDubbingTab(QWidget):
    """Scene cards, the Main Dubbing commands, and the cloned voice."""

    def __init__(self, host):
        super().__init__(host)
        self.host = host
        self.setObjectName("sleepyDubbing")
        self._editors: dict = {}
        self._cards: dict = {}
        self._preview_buttons: dict = {}
        self._dirty: dict = {}
        self._highlighters: dict = {}
        self._speed_labels: dict = {}
        self._duration_labels: dict = {}
        self._image_labels: dict = {}
        self._rates: dict = {}
        self._bookmarks: dict = {}
        self._selected_sid = None
        self._has_unsaved = False
        self._spell_lang = "it"
        self._loaded_for = ""
        self._synthesis_busy = False
        self._gpu_held = False
        self._play_sid = None
        self._worker = None
        self._thread = None
        self._full_player = None
        self._find_cursor = None
        self._build()

    def _build(self) -> None:
        root = QVBoxLayout(self)
        root.setSpacing(4)
        bar = QHBoxLayout()
        self._status = QLabel("Load scenes to start editing dubbed text.")
        self._status.setStyleSheet("color:#8E8B90; font-size:11px;")
        bar.addWidget(self._status, 1)

        load = QPushButton("Load from Scenes")
        load.setToolTip("Populate cards from scenes.yaml")
        load.clicked.connect(self.load_from_scenes)

        self._save_btn = QPushButton("💾 Save Dubbing")
        self._save_btn.setToolTip("Save all dubbed text to dubbing.yaml  (Ctrl+S)")
        self._save_btn.clicked.connect(lambda: self.save(quiet=False))

        self._all_btn = QPushButton("⏩ Dub All")
        self._all_btn.setToolTip("Synthesise the cloned voice for every scene that needs it")
        self._all_btn.clicked.connect(lambda: self.dub_all(force=False))

        self._redub_btn = QPushButton("↻ Redub All")
        self._redub_btn.setToolTip(
            "Speak every scene again with the cloned voice, "
            "even when the line has not changed (use after changing speed or language)."
        )
        self._redub_btn.clicked.connect(lambda: self.dub_all(force=True))

        self._play_btn = QPushButton("▶ Play")
        self._play_btn.setToolTip("Play all dubbed audio starting from the current scene")
        self._play_btn.clicked.connect(self.play_from_current)

        self._speed = QDoubleSpinBox()
        self._speed.setRange(0.5, 2.0)
        self._speed.setSingleStep(0.01)
        self._speed.setDecimals(2)
        self._speed.setValue(1.0)
        self._speed.setSuffix("x")
        self._speed.setFixedHeight(28)
        self._speed.setToolTip("Narration speed applied after the clone speaks (1.0 = as spoken)")

        self._voice = QComboBox()
        self._voice.setFixedHeight(28)
        self._voice.setMinimumContentsLength(16)
        for language in LANGUAGES:
            self._voice.addItem(language, language)
        self._voice.setToolTip(
            "Cloned voice language for Dub All and the scene play button. "
            "Italian and English both use your reference clip."
        )
        self._voice.currentIndexChanged.connect(self._on_voice_changed)

        export = QPushButton("📄 Export Word")
        export.setToolTip("Export all dubbed text to a .docx file in the project output folder")
        export.clicked.connect(self.export_word)

        find = QPushButton("🔍 Find")
        find.setToolTip("Open Find / Replace panel  (Ctrl+F)")
        find.setFixedHeight(28)
        find.clicked.connect(self._toggle_find)

        spell_it = QPushButton("🇮🇹 IT")
        spell_en = QPushButton("🇬🇧 EN")
        for button in (spell_it, spell_en):
            button.setFixedHeight(28)
        self._spell_buttons = {"it": spell_it, "en": spell_en}
        spell_it.clicked.connect(lambda: self.set_spell_language("it"))
        spell_en.clicked.connect(lambda: self.set_spell_language("en"))
        self._update_spell_style()

        self._duration = QLabel("00:00")
        self._duration.setToolTip("Overall length of dubbed scenes")
        self._duration.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        self._duration.setStyleSheet("color:#E6E1E5; font-size:12px; font-weight:bold;")

        self._toolbar = [load, self._save_btn, self._all_btn, self._redub_btn, self._play_btn, export, find, spell_it, spell_en]
        bar.addWidget(load)
        bar.addWidget(self._save_btn)
        bar.addWidget(self._all_btn)
        bar.addWidget(self._redub_btn)
        bar.addWidget(self._play_btn)
        bar.addWidget(QLabel("Speed:"))
        bar.addWidget(self._speed)
        bar.addWidget(QLabel("Voice:"))
        bar.addWidget(self._voice)
        bar.addWidget(export)
        bar.addWidget(find)
        bar.addWidget(spell_it)
        bar.addWidget(spell_en)
        bar.addStretch(1)
        bar.addWidget(self._duration)
        root.addLayout(bar)

        self._find_panel = QWidget()
        self._find_panel.setVisible(False)
        find_row = QHBoxLayout(self._find_panel)
        find_row.setContentsMargins(0, 2, 0, 2)
        find_row.setSpacing(4)
        self._find_input = QLineEdit()
        self._find_input.setPlaceholderText("Find…")
        self._find_input.setFixedHeight(26)
        self._find_input.returnPressed.connect(self.find_next)
        self._find_input.textChanged.connect(lambda: setattr(self, "_find_cursor", None))
        self._replace_input = QLineEdit()
        self._replace_input.setPlaceholderText("Replace with…")
        self._replace_input.setFixedHeight(26)
        self._find_case = QCheckBox("Match case")
        self._find_word = QCheckBox("Whole word")
        previous = QPushButton("▲ Prev")
        previous.setFixedHeight(26)
        previous.clicked.connect(self.find_prev)
        nxt = QPushButton("▼ Next")
        nxt.setFixedHeight(26)
        nxt.clicked.connect(self.find_next)
        replace_one = QPushButton("Replace")
        replace_one.setFixedHeight(26)
        replace_one.clicked.connect(self.replace_one)
        replace_all = QPushButton("Replace All")
        replace_all.setFixedHeight(26)
        replace_all.clicked.connect(self.replace_all)
        close_find = QPushButton("X")
        close_find.setFixedSize(26, 26)
        close_find.setToolTip("Close find/replace panel")
        close_find.setStyleSheet("QPushButton { color: #E6E1E5; font-weight: bold; }")
        close_find.clicked.connect(self._close_find)
        self._find_status = QLabel("")
        self._find_status.setStyleSheet("color:#8E8B90; font-size:11px;")
        find_row.addWidget(QLabel("Find:"))
        find_row.addWidget(self._find_input, 2)
        find_row.addWidget(QLabel("Replace:"))
        find_row.addWidget(self._replace_input, 2)
        find_row.addWidget(self._find_case)
        find_row.addWidget(self._find_word)
        find_row.addWidget(previous)
        find_row.addWidget(nxt)
        find_row.addWidget(replace_one)
        find_row.addWidget(replace_all)
        find_row.addWidget(self._find_status, 1)
        find_row.addWidget(close_find)
        root.addWidget(self._find_panel)

        find_shortcut = QShortcut(QKeySequence("Ctrl+F"), self)
        find_shortcut.setContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
        find_shortcut.activated.connect(self._open_find)
        bookmark_shortcut = QShortcut(QKeySequence("Ctrl+B"), self)
        bookmark_shortcut.setContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
        bookmark_shortcut.activated.connect(self.goto_next_bookmark)

        self._progress = QProgressBar()
        self._progress.setRange(0, 100)
        self._progress.setFixedHeight(6)
        self._progress.setTextVisible(False)
        self._progress.setStyleSheet(
            "QProgressBar { background:#0F0D13; border:none; }"
            "QProgressBar::chunk { background:#4CAF50; }"
        )
        self._progress.setVisible(False)
        root.addWidget(self._progress)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._cards_widget = QWidget()
        self._cards_layout = QVBoxLayout(self._cards_widget)
        self._cards_layout.setSpacing(6)
        self._cards_layout.addStretch(1)
        scroll.setWidget(self._cards_widget)
        self._scroll = scroll
        root.addWidget(scroll, 1)
        self._update_dub_btn()
        self._update_save_btn()

    def project(self) -> str:
        return str(getattr(self.host, "_project", "") or "").strip()

    def _output(self, name: str) -> str:
        project = self.project()
        return os.path.join(project, "output", name) if project else ""

    def _audio_path(self, scene_id: int) -> str:
        project = self.project()
        if not project:
            return ""
        return os.path.join(project, "output", "audio", f"scene_{int(scene_id):03d}.mp3")

    def voice_language(self) -> str:
        return normalize_language(str(self._voice.currentData() or "Italian"))

    def sync_language(self, language: str) -> None:
        """Follow the episode language without marking audio stale."""
        target = normalize_language(language)
        index = self._voice.findData(target)
        if index < 0 or self._voice.currentIndex() == index:
            return
        self._voice.blockSignals(True)
        self._voice.setCurrentIndex(index)
        self._voice.blockSignals(False)

    def set_busy(self, busy: bool) -> None:
        self._synthesis_busy = bool(busy)
        for button in self._toolbar:
            button.setEnabled(not busy)
        self._speed.setEnabled(not busy)
        self._voice.setEnabled(not busy)
        if not busy:
            self._update_dub_btn()
            self._update_save_btn()

    def note_project_changed(self) -> None:
        self._loaded_for = ""
        self._clear_cards()
        self._status.setText("Load scenes to start editing dubbed text.")
        self._update_dub_btn()
        self._update_save_btn()
        tabs = getattr(self.host, "inner_tabs", None)
        if tabs is not None and tabs.tabText(tabs.currentIndex()) == "Dubbing":
            self.ensure_loaded()

    def ensure_loaded(self) -> None:
        project = self.project()
        if not project or self._loaded_for == os.path.abspath(project):
            return
        scenes_path = self._output("scenes.yaml")
        if not scenes_path or not os.path.isfile(scenes_path):
            return
        self._populate_from_disk()

    def load_from_scenes(self) -> None:
        scenes_path = self._output("scenes.yaml")
        if not scenes_path or not os.path.isfile(scenes_path):
            QMessageBox.warning(self, "No scenes", "Run 'Split Scenes' first.")
            return
        scenes = (yaml.safe_load(Path(scenes_path).read_text(encoding="utf-8")) or {}).get("scenes", [])
        existing = self._read_dubbing()
        self._populate(scenes, existing, from_disk=False)

    def _populate_from_disk(self) -> None:
        scenes_path = self._output("scenes.yaml")
        if not scenes_path or not os.path.isfile(scenes_path):
            self._update_save_btn()
            self._update_dub_btn()
            return
        scenes = (yaml.safe_load(Path(scenes_path).read_text(encoding="utf-8")) or {}).get("scenes", [])
        self._populate(scenes, self._read_dubbing(), from_disk=True)

    def _read_dubbing(self) -> dict:
        path = self._output("dubbing.yaml")
        if not path or not os.path.isfile(path):
            return {}
        loaded = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
        return loaded if isinstance(loaded, dict) else {}

    def _clear_cards(self) -> None:
        while self._cards_layout.count() > 1:
            item = self._cards_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()
        self._editors = {}
        self._cards = {}
        self._preview_buttons = {}
        self._dirty = {}
        self._highlighters = {}
        self._speed_labels = {}
        self._duration_labels = {}
        self._image_labels = {}
        self._rates = {}
        self._bookmarks = {}
        self._selected_sid = None

    def _populate(self, scenes: list, existing: dict = None, from_disk: bool = False) -> None:
        existing = existing or {}
        self._clear_cards()
        bookmark_path = self._output("dubbing_bookmark.yaml")
        if bookmark_path and os.path.isfile(bookmark_path):
            try:
                data = yaml.safe_load(Path(bookmark_path).read_text(encoding="utf-8")) or {}
                if data.get("bookmarked_scene"):
                    self._bookmarks[int(data["bookmarked_scene"])] = True
            except Exception:
                pass

        for scene in scenes or []:
            if not isinstance(scene, dict) or "id" not in scene:
                continue
            try:
                scene_id = int(scene["id"])
            except (TypeError, ValueError):
                continue
            self._add_card(scene_id, str(scene.get("text") or ""), _yaml_entry(existing, scene_id))

        if self._editors:
            self._select_scene(min(self._editors))
        project = self.project()
        self._loaded_for = os.path.abspath(project) if project else ""
        self._status.setText(f"{len(self._editors)} scene(s) loaded.")
        self._update_total_duration()
        self._update_image_badges()
        self._has_unsaved = not from_disk
        self._update_save_btn()
        self._update_dub_btn()

    def _add_card(self, scene_id: int, original: str, entry: dict) -> None:
        dubbed = entry.get("dubbed")
        if not isinstance(dubbed, str) or not dubbed:
            dubbed = original
        rate = entry.get("rate_pct")
        if rate is not None:
            try:
                self._rates[scene_id] = int(rate)
            except (TypeError, ValueError):
                pass

        card = QWidget()
        card.setObjectName("dubCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(10, 8, 10, 8)
        layout.setSpacing(4)

        header = QHBoxLayout()
        title = QLabel(f"Scene {scene_id}")
        title.setStyleSheet("color:#96BDE2; font-weight:bold; font-size:11px; background:transparent; border:none;")
        chars = QLabel(f"{len(dubbed)} chars")
        chars.setStyleSheet("color:#8E8B90; font-size:10px; background:transparent; border:none;")
        chars.setAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
        speed = QLabel(speed_badge_text(self._rates.get(scene_id), os.path.isfile(self._audio_path(scene_id))))
        speed.setToolTip("Speech speed used to dub this scene")
        speed.setStyleSheet(
            "color:#C9A6E6; font-size:10px; font-weight:bold; background:#2A2438; "
            "border:1px solid #46365A; border-radius:8px; padding:1px 6px;"
        )
        duration = QLabel(self._duration_text(scene_id))
        duration.setToolTip("Length of this scene's dubbed audio")
        duration.setStyleSheet(
            "color:#9FD6B8; font-size:10px; font-weight:bold; background:#1E3A2C; "
            "border:1px solid #2F5B41; border-radius:8px; padding:1px 6px;"
        )
        images = QLabel()
        images.setStyleSheet(
            "color:#F3C98B; font-size:12px; font-weight:bold; background:#3A2D1E; "
            "border:1px solid #6B5130; border-radius:8px; padding:1px 6px;"
        )
        size_slider = QSlider(Qt.Orientation.Horizontal)
        size_slider.setRange(8, 36)
        size_slider.setValue(15)
        size_slider.setFixedWidth(90)
        size_slider.setToolTip("Adjust translated text size for this scene")
        size_title = QLabel("Text size:")
        size_title.setCursor(Qt.CursorShape.PointingHandCursor)
        size_title.setToolTip("Click to reset text size to 15px")
        size_label = QLabel("15px")
        size_label.setFixedWidth(28)
        size_label.setStyleSheet("color:#8E8B90; font-size:10px; background:transparent; border:none;")
        preview = QPushButton("▶")
        preview.setFixedSize(28, 28)
        preview.setToolTip("Synthesise and play this segment")
        preview.clicked.connect(lambda _checked=False, sid=scene_id: self.preview_segment(sid))
        deepl = QPushButton("DeepL")
        deepl.setFixedHeight(28)
        deepl.setToolTip("Select this translation and send Ctrl+C twice to DeepL")
        bookmark = QPushButton("🔖")
        bookmark.setFixedSize(28, 28)
        bookmark.setCheckable(True)
        bookmark.setChecked(bool(self._bookmarks.get(scene_id, False)))
        bookmark.setToolTip("Toggle bookmark (Ctrl+B cycles through bookmarks)")
        self._apply_bookmark_style(bookmark, bookmark.isChecked())
        bookmark.toggled.connect(
            lambda checked, sid=scene_id, button=bookmark: self._toggle_bookmark(sid, checked, button)
        )
        header.addWidget(title)
        header.addWidget(chars, 1)
        header.addWidget(speed)
        header.addWidget(duration)
        header.addWidget(images)
        header.addWidget(size_title)
        header.addWidget(size_slider)
        header.addWidget(size_label)
        header.addWidget(deepl)
        header.addWidget(bookmark)
        header.addWidget(preview)
        layout.addLayout(header)

        original_label = QLabel(original)
        original_label.setWordWrap(True)
        original_label.setStyleSheet(
            "color:#5a6470; font-size:12px; background:#131118; "
            "border:1px solid #2a2830; border-radius:2px; padding:4px;"
        )
        layout.addWidget(original_label)

        editor = QPlainTextEdit()
        editor.setPlainText(dubbed)
        editor.setPlaceholderText("Enter dubbed / translated text here…")
        editor.setMinimumHeight(70)
        editor.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        editor.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        editor.setCursorWidth(2)

        def _resize(edit=editor):
            last = edit.document().lastBlock()
            document_height = math.ceil(edit.blockBoundingGeometry(last).bottom())
            margins = edit.contentsMargins()
            chrome = margins.top() + margins.bottom() + edit.frameWidth() * 2 + 12
            target = max(70, document_height + chrome)
            if edit.height() != target:
                edit.setFixedHeight(target)

        def _on_resize(event, edit=editor, resize=_resize):
            QPlainTextEdit.resizeEvent(edit, event)
            QTimer.singleShot(0, lambda: resize(edit))

        def _apply_size(size: int, edit=editor, label=size_label, resize=_resize):
            label.setText(f"{size}px")
            edit.setStyleSheet(
                "QPlainTextEdit { background:#1D1B20; color:#E6E1E5; "
                "border:1px solid #36343B; border-radius:2px; padding:4px; "
                f"font-family:'Segoe UI',sans-serif; font-size:{int(size)}px; }}"
                "QPlainTextEdit:focus { border:1px solid #96BDE2; }"
            )
            resize(edit)

        _apply_size(size_slider.value())
        size_slider.valueChanged.connect(_apply_size)
        size_title.mousePressEvent = (
            lambda event, slider=size_slider: slider.setValue(15)
            if event.button() == Qt.MouseButton.LeftButton else None
        )
        editor.document().documentLayout().documentSizeChanged.connect(
            lambda _size, edit=editor, resize=_resize: resize(edit)
        )
        editor.resizeEvent = _on_resize

        def _on_text(sid=scene_id, label=chars, edit=editor, resize=_resize):
            label.setText(f"{len(edit.toPlainText())} chars")
            resize(edit)
            self._dirty[sid] = True
            self._has_unsaved = True
            self._apply_card_state(sid)
            self._update_save_btn()
            self._update_dub_btn()

        def _on_focus(event, sid=scene_id, edit=editor):
            self._select_scene(sid)
            QPlainTextEdit.focusInEvent(edit, event)

        editor.textChanged.connect(_on_text)
        editor.focusInEvent = _on_focus
        editor.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        editor.customContextMenuRequested.connect(
            lambda pos, edit=editor: self._context_menu(edit, pos)
        )
        deepl.clicked.connect(lambda _checked=False, edit=editor: self._send_to_deepl(edit))
        layout.addWidget(editor)

        for widget in [card, *card.findChildren(QWidget)]:
            widget.setProperty("dubSceneId", scene_id)
            widget.installEventFilter(self)

        self._editors[scene_id] = editor
        self._cards[scene_id] = card
        self._preview_buttons[scene_id] = preview
        self._speed_labels[scene_id] = speed
        self._duration_labels[scene_id] = duration
        self._image_labels[scene_id] = images
        if self._spell_lang != "off":
            editor.blockSignals(True)
            self._highlighters[scene_id] = _SpellHighlighter(editor.document(), language=self._spell_lang)
            editor.blockSignals(False)
        self._dirty[scene_id] = False
        self._cards_layout.insertWidget(self._cards_layout.count() - 1, card)
        self._apply_card_state(scene_id)

    def eventFilter(self, watched, event):
        if event.type() == QEvent.Type.MouseButtonPress:
            scene_id = watched.property("dubSceneId")
            if scene_id is not None and event.button() == Qt.MouseButton.LeftButton:
                self._select_scene(int(scene_id))
        return super().eventFilter(watched, event)

    def _select_scene(self, scene_id: int) -> None:
        previous = self._selected_sid
        self._selected_sid = scene_id
        if previous in self._cards:
            self._apply_card_state(previous)
        if scene_id in self._cards:
            self._apply_card_state(scene_id)
        self._status.setText(f"Scene {scene_id} selected for playback.")

    def _state(self, scene_id: int) -> str:
        if not os.path.isfile(self._audio_path(scene_id)):
            return "not_dubbed"
        if self._dirty.get(scene_id, False):
            return "to_redub"
        return "dubbed"

    def _apply_card_state(self, scene_id: int) -> None:
        state = self._state(scene_id)
        card = self._cards.get(scene_id)
        button = self._preview_buttons.get(scene_id)
        if card is None or button is None:
            return
        colors = {"not_dubbed": _NOT_DUBBED, "to_redub": _TO_REDUB, "dubbed": _DUBBED}[state]
        border = {"not_dubbed": "#36343B", "to_redub": "#F8B23D", "dubbed": "#4CAF50"}[state]
        if scene_id == self._selected_sid:
            border = "#96BDE2"
        card.setStyleSheet(
            f"QWidget#dubCard {{ background:#1D1B20; border:2px solid {border}; border-radius:4px; }}"
        )
        icon = {"not_dubbed": "▶", "to_redub": "⟳", "dubbed": "▶"}[state]
        button.setText(icon)
        background, foreground = colors
        button.setStyleSheet(
            f"QPushButton {{ background:{background}; color:{foreground}; border:1px solid {border}; "
            "border-radius:3px; font-size:14px; "
            "min-width:28px; max-width:28px; min-height:28px; max-height:28px; padding:0; }}"
            "QPushButton:hover { background:#2A282F; }"
        )

    def _dubbing_status(self) -> str:
        if not self._editors:
            return "none"
        missing = any(not os.path.isfile(self._audio_path(scene_id)) for scene_id in self._editors)
        if missing:
            return "all_stale"
        if any(self._dirty.get(scene_id, False) for scene_id in self._editors):
            return "stale"
        return "complete"

    def _update_dub_btn(self) -> None:
        status = self._dubbing_status()
        background, hover = _BTN_COLORS.get(status, ("#c0392b", "#a93226"))
        enabled = status != "none" and not self._synthesis_busy
        self._all_btn.setEnabled(enabled)
        self._all_btn.setStyleSheet(_BTN_SS.format(bg=background, hv=hover))
        self._redub_btn.setEnabled(enabled)

    def _update_save_btn(self) -> None:
        if not self._editors:
            background, hover, enabled = "#555555", "#444444", False
        elif self._has_unsaved:
            background, hover, enabled = "#c0392b", "#a93226", True
        else:
            background, hover, enabled = "#27ae60", "#1e8449", True
        self._save_btn.setEnabled(enabled and not self._synthesis_busy)
        self._save_btn.setStyleSheet(_BTN_SS.format(bg=background, hv=hover))

    def save(self, quiet: bool = False) -> bool:
        path = self._output("dubbing.yaml")
        if not path:
            if not quiet:
                QMessageBox.warning(self, "No project", "Load a project first.")
            return False
        if not self._editors:
            if not quiet:
                QMessageBox.warning(self, "Nothing to save", "Load scenes first.")
            return False
        originals = {}
        scenes_path = self._output("scenes.yaml")
        if scenes_path and os.path.isfile(scenes_path):
            scenes = (yaml.safe_load(Path(scenes_path).read_text(encoding="utf-8")) or {}).get("scenes", [])
            originals = {
                int(scene["id"]): str(scene.get("text") or "")
                for scene in scenes
                if isinstance(scene, dict) and "id" in scene
            }
        data = {}
        for scene_id, editor in self._editors.items():
            row = {"original": originals.get(scene_id, ""), "dubbed": editor.toPlainText()}
            if scene_id in self._rates:
                row["rate_pct"] = self._rates[scene_id]
            data[scene_id] = row
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            Path(path).write_text(
                yaml.safe_dump(data, allow_unicode=True, sort_keys=True, default_flow_style=False),
                encoding="utf-8",
            )
        except Exception as exc:
            if not quiet:
                QMessageBox.critical(self, "Save failed", str(exc))
            return False
        if not quiet:
            self._status.setText(f"Saved {len(data)} scene(s) — {path}")
        self._has_unsaved = False
        self._update_save_btn()
        self._update_dub_btn()
        return True

    def export_word(self) -> None:
        project = self.project()
        if not project:
            QMessageBox.warning(self, "No project", "Open a project first.")
            return
        if not self._editors:
            QMessageBox.warning(self, "Nothing to export", "Load scenes in the Dubbing tab first.")
            return
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
        except ImportError:
            QMessageBox.critical(
                self, "Missing dependency",
                "python-docx is not installed.\nRun: pip install python-docx",
            )
            return
        document = Document()
        document.core_properties.title = os.path.basename(project)
        style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        heading = document.add_heading(os.path.basename(project), level=1)
        heading.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        for scene_id in sorted(self._editors):
            document.add_paragraph(self._editors[scene_id].toPlainText().strip())
        out_dir = os.path.join(project, "output")
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, f"{os.path.basename(project).replace(' ', '_')}_dubbing.docx")
        try:
            document.save(out_path)
        except PermissionError:
            QMessageBox.warning(
                self, "Export failed",
                f"Could not write:\n{out_path}\n\nClose the file in Word and try again.",
            )
            return
        box = QMessageBox(self)
        box.setWindowTitle("Exported")
        box.setText(f"Word file saved to:\n{out_path}")
        box.setIcon(QMessageBox.Icon.NoIcon)
        box.exec()

    def dub_all(self, force: bool = False) -> None:
        if not self._editors:
            QMessageBox.warning(self, "No segments", "Load scenes first.")
            return
        if force:
            scene_ids = sorted(self._editors)
        else:
            scene_ids = sorted(
                scene_id for scene_id in self._editors
                if not os.path.isfile(self._audio_path(scene_id)) or self._dirty.get(scene_id, False)
            )
        if not scene_ids:
            self._status.setText("All segments already dubbed and up to date.")
            return
        spoken = []
        skipped = []
        for scene_id in scene_ids:
            text = self._spoken_text(scene_id)
            if not text:
                skipped.append(scene_id)
                continue
            spoken.append({"id": scene_id, "text": text})
        if skipped:
            listed = ", ".join(f"{scene_id:03d}" for scene_id in skipped)
            self.host._log(f"Skipped empty scenes: {listed}")
        if not spoken:
            QMessageBox.warning(self, "Nothing to speak", "The selected scenes have no dubbed text.")
            return
        self._play_sid = None
        self._start_voice(spoken, replace=True)

    def preview_segment(self, scene_id: int) -> None:
        text = self._spoken_text(scene_id)
        if not text:
            return
        self._play_sid = scene_id
        button = self._preview_buttons.get(scene_id)
        if button is not None:
            button.setEnabled(False)
            button.setText("…")
        self._start_voice([{"id": scene_id, "text": text}], replace=True)

    def _spoken_text(self, scene_id: int) -> str:
        editor = self._editors.get(scene_id)
        if editor is None:
            return ""
        text = editor.toPlainText().strip()
        if not text:
            return ""
        fixups = os.path.join(str(getattr(self.host, "_config_dir", "") or ""), "tts_fixups.yaml")
        return apply_tts_fixups(text, fixups).strip()

    def _start_voice(self, scenes: list, replace: bool) -> None:
        if not self.host._require_project() or self.host._busy:
            self._restore_preview_buttons()
            return
        self.save(quiet=True)
        language = self.voice_language()
        reference = self.host._ref_path.text().strip()
        transcript = self.host._ref_text.toPlainText().strip()
        checkpoint = self.host._checkpoint.text().strip()
        try:
            validate_request(language, reference, transcript, checkpoint or None)
        except Exception as exc:
            self.host._fail(str(exc))
            self._restore_preview_buttons()
            return
        if not self.host.controller.reserve_gpu(_GPU_OWNER):
            owner = self.host.controller.gpu_owner() or "a pipeline run"
            self.host._fail(
                f"The GPU is busy ({owner}). Wait until it is free. "
                "HiDream and the cloned voice do not run together."
            )
            self._restore_preview_buttons()
            return
        self._gpu_held = True
        from utilis.project_paths import ProjectLayout
        layout = ProjectLayout(self.project())
        layout.ensure_dirs()
        voice_run = VoiceRun()
        self.host._voice_run = voice_run
        worker = _DubWorker({
            "scenes": scenes,
            "output_dir": layout.audio,
            "timings_path": os.path.join(layout.audio, "timings.yaml"),
            "language": language,
            "ref_audio": reference,
            "ref_text": transcript,
            "model_dir": checkpoint or None,
            "seed": int(self.host._seed.value()),
            "replace_existing": replace,
        }, voice_run, float(self._speed.value()))
        thread = QThread()
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        worker.progress.connect(self._on_progress)
        worker.scene_ready.connect(self._on_scene_ready)
        worker.finished.connect(self._on_finished)
        worker.finished.connect(thread.quit)
        thread.finished.connect(self._clear_thread)
        self._worker = worker
        self._thread = thread
        self.host._voice_worker = worker
        self.host._voice_thread = thread
        self.host._set_busy(True)
        self._progress.setVisible(True)
        self._progress.setValue(0)
        self._status.setText("Loading the cloned voice…")
        self.host._log(f"Dubbing {language} with the cloned voice.")
        thread.start()

    def _on_progress(self, message: str) -> None:
        text = str(message or "").strip()
        if not text:
            return
        self._status.setText(text)
        self.host._log(text)
        parts = text.split()
        if text.startswith("Spoke scene ") and len(parts) >= 3 and "/" in parts[-1]:
            try:
                done, total = parts[-1].strip("()").split("/")
                self._progress.setValue(int(int(done) / max(int(total), 1) * 100))
            except ValueError:
                pass

    def _on_scene_ready(self, scene_id: int) -> None:
        self._dirty[scene_id] = False
        self._rates[scene_id] = rate_percent(self._speed.value())
        self._apply_card_state(scene_id)
        self._update_segment_badges(scene_id)
        self._update_total_duration()
        self._update_dub_btn()
        self.save(quiet=True)
        button = self._preview_buttons.get(scene_id)
        if button is not None:
            button.setEnabled(True)

    def _on_finished(self, ok: bool, message: str) -> None:
        if self._gpu_held:
            self.host.controller.release_gpu(_GPU_OWNER)
            self._gpu_held = False
        self.host._set_busy(False)
        self._progress.setVisible(False)
        self._restore_preview_buttons()
        play_sid = self._play_sid
        self._play_sid = None
        if ok:
            self._status.setText(message or "Dubbing finished.")
            self.host._log(message or "Dubbing finished.")
            if play_sid is not None:
                self._play_file(self._audio_path(play_sid))
            return
        if "cancel" in str(message).lower():
            self._status.setText("Dubbing cancelled")
            self.host._log(message)
            return
        self._status.setText("Dubbing failed")
        self.host._log(message)
        QMessageBox.critical(self, "Dubbing failed", message)

    def _restore_preview_buttons(self) -> None:
        for scene_id, button in self._preview_buttons.items():
            button.setEnabled(True)
            self._apply_card_state(scene_id)

    def _clear_thread(self) -> None:
        thread = self._thread
        self._worker = None
        self._thread = None
        if getattr(self.host, "_voice_thread", None) is thread:
            self.host._clear_voice()

    def stop(self) -> None:
        run = getattr(self.host, "_voice_run", None)
        if run is not None:
            run.stop()

    def _on_voice_changed(self) -> None:
        language = self.voice_language()
        if hasattr(self.host, "_language_value") and self.host._language_value() != language:
            self.host._set_language(language)
        changed = False
        for scene_id in self._editors:
            if os.path.isfile(self._audio_path(scene_id)):
                self._dirty[scene_id] = True
                self._apply_card_state(scene_id)
                changed = True
        if changed:
            self._has_unsaved = True
            self._update_dub_btn()
            self._update_save_btn()
            self._status.setText("Voice changed — re-dub segments to update audio.")

    def _duration_text(self, scene_id: int) -> str:
        path = self._audio_path(scene_id)
        if not os.path.isfile(path):
            return "--:--"
        return _clock(_audio_seconds(path))

    def _update_total_duration(self) -> None:
        total = 0.0
        for scene_id in self._editors:
            path = self._audio_path(scene_id)
            if os.path.isfile(path):
                total += _audio_seconds(path)
        self._duration.setText(_clock(total))

    def _update_segment_badges(self, scene_id: int) -> None:
        speed = self._speed_labels.get(scene_id)
        if speed is not None:
            speed.setText(speed_badge_text(
                self._rates.get(scene_id),
                os.path.isfile(self._audio_path(scene_id)),
            ))
        duration = self._duration_labels.get(scene_id)
        if duration is not None:
            duration.setText(self._duration_text(scene_id))
        self._update_image_badges(scene_id)

    def _selected_image_counts(self) -> dict:
        path = self._output("lightbox_selections.yaml")
        if not path or not os.path.isfile(path):
            return {}
        try:
            data = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
        except Exception:
            return {}
        counts = {}
        if isinstance(data, dict):
            for key, files in data.items():
                if isinstance(files, list):
                    try:
                        counts[int(key)] = len(files)
                    except (TypeError, ValueError):
                        continue
        return counts

    def _update_image_badges(self, only_sid: int = None) -> None:
        counts = self._selected_image_counts()
        scene_ids = [only_sid] if only_sid is not None else list(self._image_labels)
        for scene_id in scene_ids:
            label = self._image_labels.get(scene_id)
            if label is None:
                continue
            selected = counts.get(scene_id, 0)
            audio = _audio_seconds(self._audio_path(scene_id))
            required = required_shot_count(audio)
            label.setText(f"Images {selected} / {required}")
            if selected < required:
                background, border, color = "#3A2D1E", "#6B5130", "#F3C98B"
            elif selected == required:
                background, border, color = "#1E3A2C", "#2F5B41", "#9FD6B8"
            else:
                background, border, color = "#3A1E1E", "#6B3030", "#F3908B"
            label.setStyleSheet(
                f"color:{color}; font-size:12px; font-weight:bold; background:{background}; "
                f"border:1px solid {border}; border-radius:8px; padding:1px 6px;"
            )
            if audio <= 0:
                label.setToolTip("Dub this scene to calculate how many shots the audio needs.")
            else:
                label.setToolTip(
                    f"{selected} selected image(s); {required} required. "
                    f"{audio:.0f}s of dubbed audio at {OPTIMAL_SHOT_SECONDS:.0f}s per shot."
                )

    def _toggle_find(self) -> None:
        if self._find_panel.isVisible():
            self._close_find()
        else:
            self._open_find()

    def _open_find(self) -> None:
        self._find_panel.setVisible(True)
        self._find_input.setFocus()
        self._find_input.selectAll()
        self._find_cursor = None

    def _close_find(self) -> None:
        self._find_panel.setVisible(False)
        self._find_status.setText("")
        self._find_cursor = None

    def _editors_in_order(self):
        return sorted(self._editors.items())

    def _search_flags(self):
        return 0 if self._find_case.isChecked() else re.IGNORECASE

    def _find_spans(self, text: str, query: str) -> list:
        if not query:
            return []
        pattern = re.escape(query)
        if self._find_word.isChecked():
            pattern = r"\b" + pattern + r"\b"
        return [(match.start(), match.end()) for match in re.finditer(pattern, text, self._search_flags())]

    def find_next(self) -> None:
        self._find_step(True)

    def find_prev(self) -> None:
        self._find_step(False)

    def _find_step(self, forward: bool) -> None:
        query = self._find_input.text()
        if not query:
            return
        matches = []
        for scene_id, editor in self._editors_in_order():
            for start, end in self._find_spans(editor.toPlainText(), query):
                matches.append((scene_id, editor, start, end))
        if not matches:
            self._find_status.setText("No matches")
            self._find_cursor = None
            return
        self._find_status.setText(f"{len(matches)} match(es)")
        cursor = self._find_cursor
        if cursor is None:
            index = 0 if forward else len(matches) - 1
        elif forward:
            index = (cursor + 1) % len(matches)
        else:
            index = (cursor - 1) % len(matches)
        self._find_cursor = index
        _scene_id, editor, start, end = matches[index]
        self._highlight(editor, start, end)

    def _highlight(self, editor, start: int, end: int) -> None:
        cursor = editor.textCursor()
        cursor.setPosition(start)
        cursor.setPosition(end, QTextCursor.MoveMode.KeepAnchor)
        editor.setTextCursor(cursor)
        editor.setFocus()
        editor.ensureCursorVisible()
        card = editor.parentWidget()
        if card is not None:
            self._scroll.ensureWidgetVisible(card)

    def replace_one(self) -> None:
        query = self._find_input.text()
        replacement = self._replace_input.text()
        if not query:
            return
        for scene_id, editor in self._editors_in_order():
            if not editor.hasFocus():
                continue
            cursor = editor.textCursor()
            if cursor.hasSelection():
                pattern = re.escape(query)
                if self._find_word.isChecked():
                    pattern = r"\b" + pattern + r"\b"
                if re.fullmatch(pattern, cursor.selectedText(), self._search_flags()):
                    cursor.insertText(replacement)
                    self._dirty[scene_id] = True
                    self._has_unsaved = True
                    self._update_dub_btn()
                    self._update_save_btn()
            break
        self.find_next()

    def replace_all(self) -> None:
        query = self._find_input.text()
        replacement = self._replace_input.text()
        if not query:
            return
        pattern = re.escape(query)
        if self._find_word.isChecked():
            pattern = r"\b" + pattern + r"\b"
        count = 0
        for scene_id, editor in self._editors_in_order():
            updated, found = re.subn(pattern, replacement, editor.toPlainText(), flags=self._search_flags())
            if found:
                editor.setPlainText(updated)
                self._dirty[scene_id] = True
                self._has_unsaved = True
                count += found
        self._update_dub_btn()
        self._update_save_btn()
        self._find_status.setText(f"Replaced {count} occurrence(s)")

    def set_spell_language(self, language: str) -> None:
        self._spell_lang = language
        self._update_spell_style()
        for scene_id, editor in self._editors.items():
            editor.blockSignals(True)
            highlighter = self._highlighters.get(scene_id)
            if highlighter is not None:
                highlighter.set_language(language)
            else:
                self._highlighters[scene_id] = _SpellHighlighter(editor.document(), language=language)
            editor.blockSignals(False)

    def _update_spell_style(self) -> None:
        active = (
            "QPushButton { background:#2A5298; color:#FFFFFF; border:1px solid #4A72B8; "
            "border-radius:3px; font-size:12px; padding:0 6px; }"
            "QPushButton:hover { background:#3A63A8; }"
        )
        inactive = (
            "QPushButton { background:#2A2830; color:#8E8B90; border:1px solid #3A3840; "
            "border-radius:3px; font-size:12px; padding:0 6px; }"
            "QPushButton:hover { background:#3A3640; }"
        )
        for key, button in self._spell_buttons.items():
            button.setStyleSheet(active if key == self._spell_lang else inactive)

    def _reload_custom_words(self) -> None:
        _SpellHighlighter._custom_words = _SpellHighlighter._load_custom_words()
        for scene_id, highlighter in self._highlighters.items():
            if highlighter._checker is not None and _SpellHighlighter._custom_words:
                highlighter._checker.word_frequency.load_words(_SpellHighlighter._custom_words)
            editor = self._editors.get(scene_id)
            if editor is not None:
                editor.blockSignals(True)
            try:
                highlighter.rehighlight()
            finally:
                if editor is not None:
                    editor.blockSignals(False)
        self._status.setText(f"Custom word list reloaded — {len(_SpellHighlighter._custom_words)} word(s).")

    def _context_menu(self, editor: QPlainTextEdit, pos) -> None:
        cursor = editor.textCursor()
        if not cursor.hasSelection():
            cursor = editor.cursorForPosition(pos)
            cursor.select(QTextCursor.SelectionType.WordUnderCursor)
        word = cursor.selectedText().strip()
        menu = editor.createStandardContextMenu()
        menu.addSeparator()
        add_action = menu.addAction(f'Add "{word}" to Dictionary' if word else "Add to Dictionary")
        add_action.setEnabled(
            self._spell_lang == "it"
            and bool(re.fullmatch(r"[A-Za-zÀ-ÖØ-öø-ÿ']+", word))
            and word.casefold() not in _SpellHighlighter._custom_words
        )
        chosen = menu.exec(editor.viewport().mapToGlobal(pos))
        if chosen != add_action:
            return
        try:
            added = _SpellHighlighter.add_custom_word(word)
        except OSError as exc:
            QMessageBox.warning(self, "Dictionary", f"Could not update the dictionary:\n{exc}")
            return
        self._reload_custom_words()
        if added:
            self._status.setText(f'Added "{word}" to the custom dictionary.')

    def _send_to_deepl(self, editor: QPlainTextEdit) -> None:
        import ctypes
        editor.setFocus()
        editor.selectAll()
        user32 = ctypes.windll.user32
        key_up = 0x0002
        for _ in range(2):
            user32.keybd_event(0x11, 0, 0, 0)
            user32.keybd_event(0x43, 0, 0, 0)
            user32.keybd_event(0x43, 0, key_up, 0)
            user32.keybd_event(0x11, 0, key_up, 0)

    def _apply_bookmark_style(self, button, active: bool) -> None:
        if active:
            button.setStyleSheet(
                "QPushButton { background:#4A3800; color:#FFD600; border:1px solid #FFD600; "
                "border-radius:3px; font-size:14px; min-width:28px; min-height:28px; padding:0; }"
                "QPushButton:hover { background:#5A4A00; }"
            )
        else:
            button.setStyleSheet(
                "QPushButton { background:#2A2830; color:#8E8B90; border:1px solid #36343B; "
                "border-radius:3px; font-size:14px; min-width:28px; min-height:28px; padding:0; }"
                "QPushButton:hover { background:#36343B; }"
            )

    def _toggle_bookmark(self, scene_id: int, checked: bool, button) -> None:
        if checked:
            for other_id, active in list(self._bookmarks.items()):
                if active and other_id != scene_id:
                    self._bookmarks[other_id] = False
                    card = self._cards.get(other_id)
                    if card is None:
                        continue
                    for child in card.findChildren(QPushButton):
                        if child.isCheckable() and child.text() == "🔖":
                            child.blockSignals(True)
                            child.setChecked(False)
                            child.blockSignals(False)
                            self._apply_bookmark_style(child, False)
        self._bookmarks[scene_id] = checked
        self._apply_bookmark_style(button, checked)
        path = self._output("dubbing_bookmark.yaml")
        if not path:
            return
        os.makedirs(os.path.dirname(path), exist_ok=True)
        Path(path).write_text(
            yaml.safe_dump({"bookmarked_scene": scene_id if checked else None}, allow_unicode=True),
            encoding="utf-8",
        )

    def goto_next_bookmark(self) -> None:
        marked = [scene_id for scene_id, active in self._bookmarks.items() if active]
        if not marked:
            return
        card = self._cards.get(marked[0])
        if card is not None:
            self._scroll.ensureWidgetVisible(card)
            self._select_scene(marked[0])

    def play_from_current(self) -> None:
        if not self._editors:
            return
        scene_ids = sorted(self._editors)
        start = self._selected_sid or scene_ids[0]
        queue = [
            self._audio_path(scene_id)
            for scene_id in scene_ids
            if scene_id >= start and os.path.isfile(self._audio_path(scene_id))
        ]
        if not queue:
            self._status.setText("No audio files found from this scene onward.")
            return
        self._play_queue = queue
        self._play_index = 0
        self._play_btn.setText("⏹ Stop")
        try:
            self._play_btn.clicked.disconnect()
        except Exception:
            pass
        self._play_btn.clicked.connect(self.stop_playback)
        self._advance_playback()

    def _advance_playback(self) -> None:
        queue = getattr(self, "_play_queue", [])
        index = getattr(self, "_play_index", 0)
        if index >= len(queue):
            self.stop_playback()
            return
        path = queue[index]
        self._play_index = index + 1
        try:
            scene_id = int(os.path.basename(path).split("_")[1].split(".")[0])
            card = self._cards.get(scene_id)
            self._select_scene(scene_id)
            if card is not None:
                self._scroll.ensureWidgetVisible(card, 0, 12)
            self._status.setText(f"Playing scene {scene_id}  ({index + 1}/{len(queue)})")
        except Exception:
            pass
        self._play_file(path, chain=True)

    def _play_file(self, path: str, chain: bool = False) -> None:
        if not path or not os.path.isfile(path):
            return
        if self._full_player is None:
            self._full_player = QMediaPlayer(self)
            self._full_audio = QAudioOutput(self)
            self._full_player.setAudioOutput(self._full_audio)
            self._full_audio.setVolume(1.0)
        try:
            self._full_player.playbackStateChanged.disconnect(self._on_playback_state)
        except Exception:
            pass
        if chain:
            self._full_player.playbackStateChanged.connect(self._on_playback_state)
        self._full_player.setSource(QUrl.fromLocalFile(os.path.abspath(path)))
        self._full_player.play()

    def _on_playback_state(self, state) -> None:
        if state == QMediaPlayer.PlaybackState.StoppedState:
            self._advance_playback()

    def stop_playback(self) -> None:
        if self._full_player is not None:
            try:
                self._full_player.playbackStateChanged.disconnect(self._on_playback_state)
            except Exception:
                pass
            self._full_player.stop()
        self._play_queue = []
        self._play_btn.setText("▶ Play")
        try:
            self._play_btn.clicked.disconnect()
        except Exception:
            pass
        self._play_btn.clicked.connect(self.play_from_current)
        self._status.setText("Playback stopped.")
