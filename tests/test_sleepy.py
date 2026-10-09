"""Sleepy episode contract: one HiDream seed, storybook prompts, clone voice guard."""

import ast
import os
import sys
import tempfile
import unittest
from pathlib import Path

from audio_loudness import configured_loudness
from narration import qwen_voice
from narration.qwen_voice import (
    chunks_for_scene,
    chunk_narration,
    validate_request,
    voice_python,
    default_model_dir,
    worker_script,
)
from sleepy.chapter import (
    add_narration,
    align_narration_to_scenes,
    align_stills,
    build_model_prompts,
    create_episode,
    episode_folder_name,
    ensure_rome_softly_prompt,
    filenames_for_offsets,
    image_prompt_map,
    is_sleepy_project,
    lightbox_cards,
    load_episode_settings,
    load_scenes,
    load_stills,
    narration_path,
    parse_image_script,
    prompts_missing,
    publish_scenes,
    save_stills,
    selections_for_stills,
    sleepy_pipeline_config,
    split_scenes,
    still_file_path,
    list_lightbox_stills,
)


def _narration_docx(path: str, blocks) -> None:
    """Write a small .docx. Each block is (text, style name or "")."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    if "Titolo 1" not in [style.name for style in document.styles]:
        document.styles.add_style("Titolo 1", WD_STYLE_TYPE.PARAGRAPH)
    for text, style in blocks:
        paragraph = document.add_paragraph(text)
        if style:
            paragraph.style = style
    document.add_table(rows=1, cols=2)
    document.tables[0].cell(0, 0).text = "DO NOT READ TABLE LEFT"
    document.tables[0].cell(0, 1).text = "DO NOT READ TABLE RIGHT"
    document.save(path)


class SeedContractTests(unittest.TestCase):
    def test_main_keeps_three_seed_neighbours(self):
        names = filenames_for_offsets(4, 2, None)
        self.assertEqual(names, [
            "scene_004_hidream_b02_v1.png",
            "scene_004_hidream_b02_v2.png",
            "scene_004_hidream_b02_v3.png",
        ])

    def test_sleepy_uses_one_seed_named_v1(self):
        names = filenames_for_offsets(4, 2, [0])
        self.assertEqual(names, ["scene_004_hidream_b02_v1.png"])

    def test_pipeline_asks_for_offsets_instead_of_a_fixed_triple(self):
        source = Path("ui/pipeline_controller.py").read_text(encoding="utf-8")
        self.assertIn("resolve_seed_offsets", source)
        self.assertIn("variant_filename", source)
        self.assertNotIn("seed_offsets = [-1, 0, 1]", source)

    def test_main_loudness_default_stays_and_sleepy_uses_the_series_target(self):
        self.assertEqual(configured_loudness({}), (-17.0, -1.0))
        config = sleepy_pipeline_config(7, "rome_softly_ch01")
        self.assertEqual(configured_loudness(config), (-20.0, -3.0))
        self.assertEqual(config["seed_offsets"], [0])
        self.assertEqual(config["enabled_image_models"], ["hidream"])
        self.assertEqual(config["visual_style"], "rome_softly")
        self.assertEqual((config["image_width"], config["image_height"]), (1344, 768))
        self.assertEqual(config["lightbox_model_key"], "")
        self.assertEqual(config["ken_burns_motion"], "sleepy")
        self.assertNotIn("schnell", config["enabled_image_models"])


class ChapterTests(unittest.TestCase):
    def test_storybook_sentence_is_kept_and_not_turned_into_a_photograph(self):
        prompt = ensure_rome_softly_prompt(
            "Wide shot of the Tiber at dawn in 509 BC, distant figures on the bank."
        )
        self.assertIn("flat vector-style illustration", prompt.lower())
        self.assertIn("tuff stone grey, roman ochre, terracotta and muted bronze", prompt.lower())
        self.assertIn("509 BC", prompt)
        self.assertNotIn("gouache", prompt.lower())
        self.assertNotIn("sharp photograph", prompt.lower())
        self.assertNotIn("sharp focus, clear air", prompt.lower())
        again = ensure_rome_softly_prompt(prompt)
        self.assertEqual(again, prompt)
        night = (
            "A moonlit bend of the Tiber, Rome about 509 BC, no people. "
            "Style: flat vector-style illustration, clean simplified shapes, "
            "limited muted palette of deep indigo and dusk blue with a single warm light source, "
            "no text, no gore."
        )
        self.assertEqual(ensure_rome_softly_prompt(night), " ".join(night.split()))
        legacy = (
            "Wide view of the Tiber. Painterly storybook illustration in gouache "
            "and watercolour on textured paper, no text."
        )
        self.assertEqual(ensure_rome_softly_prompt(legacy), legacy)

    def test_episode_roundtrip_keeps_one_prompt_per_still(self):
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Roma capitolo 1")
            self.assertTrue(is_sleepy_project(project))
            text = "Prima scena sul Tevere all'alba.\n\nSeconda scena dentro la capanna."
            scenes = publish_scenes(project, "Italian", text)
            self.assertEqual([scene["id"] for scene in scenes], [1, 2])
            stills = align_stills(scenes, [{
                "scene_id": 1,
                "beat": 1,
                "prompt": "Dawn on the river.",
            }, {
                "scene_id": 1,
                "beat": 2,
                "prompt": "The hut interior.",
            }])
            self.assertEqual([(row["scene_id"], row["beat"]) for row in stills], [
                (1, 1), (1, 2), (2, 1),
            ])
            self.assertEqual(prompts_missing(stills), ["Scene 002 still 1"])
            stills[2]["prompt"] = "A timber hut."
            prompts = build_model_prompts(scenes, stills)
            self.assertEqual(set(prompts), {1, 2})
            row = prompts[1]["models"]["hidream"]["prompts"][0]
            self.assertEqual(row["source"], "manually_edited")
            self.assertIn("flat vector-style illustration", row["text"].lower())
            self.assertIn("tuff stone grey", row["text"].lower())
            self.assertNotIn("dev", prompts[1]["models"])
            selections, missing = selections_for_stills(
                stills, os.path.join(project, "output", "lightbox"),
            )
            self.assertEqual(selections, {})
            self.assertIn("scene_001_hidream_b01_v1.png", missing)
            self.assertIn("scene_001_hidream_b02_v1.png", missing)

    def test_a_main_project_is_not_a_sleepy_episode(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertFalse(is_sleepy_project(folder))

    def test_word_intake_keeps_both_languages_and_splits_the_italian(self):
        with tempfile.TemporaryDirectory() as folder:
            italian = os.path.join(folder, "narration_it.docx")
            english = os.path.join(folder, "narration_en.docx")
            _narration_docx(italian, [
                ("Roma, sottovoce", "Title"),
                ("Capitolo 1", "Heading 1"),
                ("Da non leggere", "Titolo 1"),
                ("", ""),
                ("Il  Tevere\xa0scorre piano davanti alle capanne all'alba.", ""),
                ("Dentro la capanna il fuoco è basso e nessuno parla.", ""),
            ])
            _narration_docx(english, [
                ("Chapter 1", "Heading 1"),
                ("The Tiber moves slowly past the huts at dawn.", ""),
                ("Inside the hut the fire is low and nobody speaks.", ""),
            ])
            project = create_episode(
                folder, "Episodio 1",
                italian_docx=italian,
                english_docx=english,
            )
            italian_text = Path(narration_path(project, "Italian")).read_text(encoding="utf-8")
            english_text = Path(narration_path(project, "English")).read_text(encoding="utf-8")
            self.assertEqual(
                italian_text,
                "Il Tevere scorre piano davanti alle capanne all'alba.\n\n"
                "Dentro la capanna il fuoco è basso e nessuno parla.\n",
            )
            self.assertEqual(
                english_text,
                "The Tiber moves slowly past the huts at dawn.\n\n"
                "Inside the hut the fire is low and nobody speaks.\n",
            )
            self.assertNotIn("DO NOT READ TABLE", italian_text)
            self.assertNotIn("Capitolo", italian_text)
            scenes = load_scenes(project)
            self.assertEqual([scene["id"] for scene in scenes], [1, 2])
            self.assertIn("Tevere", scenes[0]["text"])
            self.assertNotIn("Tiber", scenes[0]["text"])
            mirrored = Path(project, "input", "narration.txt").read_text(encoding="utf-8")
            self.assertEqual(mirrored, italian_text)
            self.assertTrue(os.path.isfile(os.path.join(project, "input", "narration_it.docx")))
            self.assertTrue(os.path.isfile(os.path.join(project, "input", "narration_en.docx")))
            settings = load_episode_settings(project)
            self.assertEqual(settings["language"], "Italian")
            self.assertEqual(settings["scenes_language"], "Italian")
            self.assertTrue(is_sleepy_project(project))

    def test_one_still_can_be_painted_without_the_others(self):
        config = sleepy_pipeline_config(42, "rome_softly_ch01", replace=True, scene_id=4, beat=2)
        self.assertEqual(config["lightbox_scene_id"], 4)
        self.assertEqual(config["lightbox_beat_index"], 2)
        self.assertEqual(config["lightbox_model_key"], "hidream")
        self.assertTrue(config["force_lightbox_update"])
        whole = sleepy_pipeline_config(42, "rome_softly_ch01")
        self.assertEqual(whole["lightbox_scene_id"], 0)
        self.assertEqual(whole["lightbox_model_key"], "")
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Solo uno")
            lightbox = os.path.join(project, "output", "lightbox")
            os.makedirs(lightbox, exist_ok=True)
            name = "scene_004_hidream_b02_v1.png"
            Path(os.path.join(lightbox, name)).write_bytes(b"png")
            rows = list_lightbox_stills(project)
            self.assertEqual(
                [(row["scene_id"], row["beat"], row["filename"]) for row in rows],
                [(4, 2, name)],
            )

    def test_still_file_path_is_the_single_final_image(self):
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Anteprima")
            path = still_file_path(project, 3, 2)
            self.assertTrue(path.endswith(os.path.join(
                "output", "lightbox", "scene_003_hidream_b02_v1.png",
            )))

    def test_image_cues_are_one_scene_and_the_prompt_is_not_spoken(self):
        script = (
            "Note di produzione (non leggere): questa riga non è una scena.\n\n"
            "[IMG 01] [IMMAGINE: Il Tevere al crepuscolo. Testo sobrio: \"Roma\".]\n\n"
            "IMG 01 — PROMPT (EN): Wide view of the Tiber at dusk, distant huts. "
            "Style: painterly storybook illustration in gouache and watercolour, no text.\n\n"
            "Buonasera, e benvenuti.\n\n"
            "Se siete arrivati fin qui, la giornata è stata lunga.\n\n"
            "[IMG 02] [IMMAGINE: Uno studio a lume di lampada.]\n\n"
            "IMG 02 — PROMPT (EN): A quiet study at night, one lamp, no people. "
            "Style: painterly storybook illustration, no text.\n\n"
            "In questa serie la racconteremo con calma.\n\n"
            "[IMG 03] [IMMAGINE: Cielo stellato. Fine.]\n\n"
            "IMG 03 — PROMPT (EN): A starry sky over a field. "
            "Style: painterly storybook illustration, no text.\n\n"
            "Tito Livio, Ab Urbe condita, praefatio; I.56–60 (Lucrezia); II.1–40 (Porsenna).\n\n"
            "Dionigi di Alicarnasso, Antichità romane, libri IV–XI "
            "(in particolare VI.13 sui Dioscuri, VI.95 sul foedus).\n\n"
            "T. J. Cornell, The Beginnings of Rome, Routledge 1995.\n"
        )
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Novantasette")
            scenes = publish_scenes(project, "Italian", script, stills=[{
                "scene_id": 1,
                "beat": 1,
                "prompt": "The prompt I typed by hand.",
            }])
            self.assertEqual([scene["id"] for scene in scenes], [1, 2, 3])
            self.assertIn("Buonasera", scenes[0]["text"])
            self.assertIn("giornata è stata lunga", scenes[0]["text"])
            self.assertNotIn("PROMPT", scenes[0]["text"])
            self.assertNotIn("IMMAGINE", scenes[0]["text"])
            self.assertNotIn("non leggere", scenes[0]["text"].lower())
            self.assertEqual(
                scenes[0]["caption"],
                'Il Tevere al crepuscolo. Testo sobrio: "Roma".',
            )
            self.assertEqual(scenes[1]["text"], "In questa serie la racconteremo con calma.")
            self.assertEqual(scenes[2]["text"], "")
            self.assertNotIn("Routledge", "\n".join(scene["text"] for scene in scenes))
            stills = load_stills(project)
            by_scene = {int(row["scene_id"]): row["prompt"] for row in stills}
            self.assertEqual(by_scene[1], "The prompt I typed by hand.")
            self.assertIn("quiet study", by_scene[2])
            self.assertIn("starry sky", by_scene[3])
            self.assertEqual(
                [(row["scene_id"], row["beat"]) for row in stills],
                [(1, 1), (2, 1), (3, 1)],
            )

    def test_palette_tag_on_the_prompt_line_is_not_spoken(self):
        script = (
            "[IMG 01] [IMMAGINE: Il Tevere al crepuscolo.]\n\n"
            "IMG 01 — PROMPT (EN) [palette: Rural / domestic / dusk]: "
            "Wide view of the Tiber at dusk, distant huts. "
            "Style: flat vector-style illustration, no text.\n\n"
            "Buonasera, e benvenuti.\n\n"
            "[IMG 02] [IMMAGINE: Cielo stellato.]\n\n"
            "IMG 02 — PROMPT (EN): A starry sky over a field.\n\n"
            "Fine della storia.\n"
        )
        scenes = split_scenes(script)
        self.assertEqual([scene["id"] for scene in scenes], [1, 2])
        self.assertEqual(scenes[0]["text"], "Buonasera, e benvenuti.")
        self.assertEqual(scenes[1]["text"], "Fine della storia.")
        self.assertNotIn("PROMPT", scenes[0]["text"])
        self.assertNotIn("palette", scenes[0]["text"].lower())
        prompts = image_prompt_map(script)
        self.assertTrue(
            prompts[1].startswith("[palette: Rural / domestic / dusk]: Wide view")
        )
        self.assertEqual(prompts[2], "A starry sky over a field.")

    def test_benvenuti_starts_speech_and_immagine_is_the_caption(self):
        script = (
            "Note di produzione (non leggere): non è una scena.\n\n"
            "Titolo da non leggere prima del benvenuto.\n\n"
            "[IMG 01] [IMMAGINE: Il Tevere al crepuscolo, con le capanne lontane.]\n\n"
            "IMG 01 — PROMPT (EN) [palette: Rural / domestic / dusk]: Wide view of the Tiber.\n\n"
            "Questa frase sta prima del benvenuto e non si legge.\n\n"
            "Buonasera, e benvenuti. Qui inizia la narrazione.\n\n"
            "[IMG 02] [IMMAGINE: Uno studio a lume di lampada.]\n\n"
            "IMG 02 — PROMPT (EN): A quiet study at night.\n\n"
            "La seconda scena continua da qui.\n"
        )
        parsed = parse_image_script(script)
        self.assertEqual(
            parsed[0]["caption"],
            "Il Tevere al crepuscolo, con le capanne lontane.",
        )
        self.assertEqual(parsed[1]["caption"], "Uno studio a lume di lampada.")
        self.assertIn("benvenuti", parsed[0]["text"].lower())
        self.assertIn("Qui inizia", parsed[0]["text"])
        self.assertNotIn("prima del benvenuto", parsed[0]["text"].lower())
        self.assertNotIn("Titolo da non leggere", parsed[0]["text"])
        self.assertNotIn("PROMPT", parsed[0]["text"])
        self.assertNotIn("Wide view", parsed[0]["text"])
        self.assertEqual(parsed[1]["text"], "La seconda scena continua da qui.")
        self.assertTrue(parsed[0]["prompt"].startswith("[palette: Rural"))

    def test_speech_is_kept_when_the_script_has_no_welcome(self):
        script = (
            "[IMG 01] [IMMAGINE: Il fiume.]\n\n"
            "IMG 01 — PROMPT (EN): A river at dusk.\n\n"
            "Il racconto comincia senza la formula di benvenuto.\n"
        )
        parsed = parse_image_script(script)
        self.assertEqual(parsed[0]["caption"], "Il fiume.")
        self.assertIn("racconto", parsed[0]["text"])
        self.assertNotIn("A river", parsed[0]["text"])

    def test_lightbox_cards_skip_a_deleted_still_file(self):
        script = (
            "[IMG 01] [IMMAGINE: Il Tevere al crepuscolo.]\n\n"
            "IMG 01 — PROMPT (EN): Wide view of the Tiber.\n\n"
            "Buonasera, e benvenuti.\n\n"
            "[IMG 02] [IMMAGINE: Uno studio a lume di lampada.]\n\n"
            "IMG 02 — PROMPT (EN): A quiet study at night.\n\n"
            "La seconda scena.\n"
        )
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Cartelle")
            publish_scenes(project, "Italian", script)
            lightbox = os.path.join(project, "output", "lightbox")
            os.makedirs(lightbox, exist_ok=True)
            for name in (
                "scene_001_hidream_b01_v1.png",
                "scene_001_hidream_b02_v1.png",
            ):
                Path(os.path.join(lightbox, name)).write_bytes(b"png")
            cards = lightbox_cards(load_scenes(project), load_stills(project), project)
            self.assertEqual([card["scene_id"] for card in cards], [1, 2])
            self.assertEqual(cards[0]["caption"], "Il Tevere al crepuscolo.")
            self.assertEqual(cards[1]["caption"], "Uno studio a lume di lampada.")
            self.assertEqual(
                [os.path.basename(image["path"]) for image in cards[0]["images"]],
                ["scene_001_hidream_b01_v1.png"],
            )
            self.assertEqual(cards[1]["images"], [{"beat": 1, "path": ""}])

    def test_a_colon_in_the_title_is_saved_as_a_folder_name(self):
        self.assertEqual(episode_folder_name("Ep. 01: Il Tevere"), "Ep. 01 - Il Tevere")
        self.assertEqual(episode_folder_name('Roma "sottovoce"'), "Roma sottovoce")
        self.assertEqual(episode_folder_name("Prima l'italiano"), "Prima l'italiano")
        with self.assertRaises(ValueError) as raised:
            episode_folder_name("   ")
        self.assertIn("episode name", str(raised.exception).lower())
        with tempfile.TemporaryDirectory() as folder:
            project = create_episode(folder, "Ep. 01: Il Tevere")
            self.assertTrue(project.endswith(os.path.join("Ep. 01 - Il Tevere")))
            self.assertTrue(is_sleepy_project(project))

    def test_chapter_lock_is_stored_on_the_episode_that_was_created(self):
        with tempfile.TemporaryDirectory() as folder:
            italian = os.path.join(folder, "it.docx")
            _narration_docx(italian, [
                ("Prima scena sul Tevere all'alba, con le capanne lontane.", ""),
            ])
            first = create_episode(
                folder, "Uno",
                italian_docx=italian,
                profile_key="rome_softly_ch01",
            )
            second = create_episode(folder, "Due", profile_key="rome_softly_ch02")
            self.assertEqual(load_episode_settings(first)["profile_key"], "rome_softly_ch01")
            self.assertEqual(load_episode_settings(second)["profile_key"], "rome_softly_ch02")
            self.assertEqual(load_scenes(second), [])
            self.assertEqual(Path(narration_path(first, "English")).read_text(encoding="utf-8"), "")
            self.assertIn("Tevere", load_scenes(first)[0]["text"])
            self.assertEqual(
                sleepy_pipeline_config(1, "rome_softly_ch02")["project_profile_key"],
                "rome_softly_ch02",
            )

    def test_english_word_can_be_added_after_the_italian_episode(self):
        with tempfile.TemporaryDirectory() as folder:
            italian = os.path.join(folder, "it.docx")
            _narration_docx(italian, [
                ("Capitolo 1", "Heading 1"),
                ("Prima scena sul Tevere all'alba, con le capanne lontane.", ""),
                ("Seconda scena dentro la capanna, con il fuoco basso.", ""),
            ])
            project = create_episode(folder, "Prima l'italiano", italian_docx=italian)
            self.assertEqual(Path(narration_path(project, "English")).read_text(encoding="utf-8"), "")
            self.assertEqual(load_episode_settings(project)["scenes_language"], "Italian")
            scenes_before = load_scenes(project)
            self.assertEqual([scene["id"] for scene in scenes_before], [1, 2])
            save_stills(project, [{
                "scene_id": 1,
                "beat": 1,
                "prompt": "Dawn on the river.",
            }])
            narration_txt = Path(project, "input", "narration.txt").read_text(encoding="utf-8")

            short_english = os.path.join(folder, "en_short.docx")
            _narration_docx(short_english, [
                ("Chapter 1", "Heading 1"),
                ("The Tiber moves slowly past the huts at dawn.", ""),
            ])
            with self.assertRaises(ValueError):
                add_narration(folder, "English", short_english)
            stored = add_narration(project, "English", short_english)
            self.assertIn("Tiber", stored)
            self.assertEqual(load_scenes(project), scenes_before)
            self.assertEqual(
                Path(project, "input", "narration.txt").read_text(encoding="utf-8"),
                narration_txt,
            )
            self.assertEqual(load_episode_settings(project)["scenes_language"], "Italian")
            self.assertEqual(load_stills(project)[0]["prompt"], "Dawn on the river.")
            self.assertIn("Prima scena", Path(narration_path(project, "Italian")).read_text(encoding="utf-8"))
            self.assertTrue(os.path.isfile(os.path.join(project, "input", "narration_en.docx")))

            revised = os.path.join(folder, "it_revised.docx")
            _narration_docx(revised, [
                ("Capitolo 1", "Heading 1"),
                ("Il Tevere di notte scorre piano sotto la luna piena.", ""),
                ("Nella capanna il fuoco è spento e la voce è più bassa.", ""),
                ("Una terza scena arriva solo nel copione rivisto, con calma.", ""),
            ])
            updated = add_narration(project, "Italian", revised)
            self.assertIn("luna piena", updated)
            self.assertNotIn("Capitolo", updated)
            self.assertEqual(load_scenes(project), scenes_before)
            self.assertEqual(
                Path(project, "input", "narration.txt").read_text(encoding="utf-8"),
                narration_txt,
            )
            self.assertEqual(load_stills(project)[0]["prompt"], "Dawn on the river.")
            self.assertIn("Tiber", Path(narration_path(project, "English")).read_text(encoding="utf-8"))
            self.assertIn("luna piena", Path(narration_path(project, "Italian")).read_text(encoding="utf-8"))
            self.assertTrue(os.path.isfile(os.path.join(project, "input", "narration_it.docx")))
            with self.assertRaises(ValueError) as raised:
                align_narration_to_scenes(stored, scenes_before, "English")
            self.assertIn("picture scenes", str(raised.exception))

            matching = os.path.join(folder, "en_full.docx")
            _narration_docx(matching, [
                ("The Tiber moves slowly past the huts at dawn.", ""),
                ("Inside the hut the fire is low and nobody speaks.", ""),
            ])
            matching_text = add_narration(project, "English", matching)
            self.assertEqual(load_scenes(project), scenes_before)
            spoken = align_narration_to_scenes(matching_text, scenes_before, "English")
            self.assertEqual([row["id"] for row in spoken], [1, 2])
            self.assertIn("Tiber", spoken[0]["text"])
            self.assertIn("Tevere", load_scenes(project)[0]["text"])

            legacy = os.path.join(folder, "late.doc")
            Path(legacy).write_bytes(b"not a docx")
            with self.assertRaises(ValueError):
                add_narration(project, "English", legacy)
            self.assertIn(
                "Inside the hut",
                Path(narration_path(project, "English")).read_text(encoding="utf-8"),
            )

    def test_english_only_word_file_becomes_the_split_language(self):
        with tempfile.TemporaryDirectory() as folder:
            english = os.path.join(folder, "rome.docx")
            _narration_docx(english, [
                ("The river is quiet and the bank is empty at dawn.", ""),
            ])
            project = create_episode(folder, "English only", english_docx=english)
            self.assertEqual(Path(narration_path(project, "Italian")).read_text(encoding="utf-8"), "")
            scenes = load_scenes(project)
            self.assertEqual(len(scenes), 1)
            self.assertIn("river", scenes[0]["text"])
            self.assertEqual(load_episode_settings(project)["language"], "English")
            self.assertEqual(load_episode_settings(project)["scenes_language"], "English")

    def test_unequal_word_files_still_create_the_episode(self):
        with tempfile.TemporaryDirectory() as folder:
            italian = os.path.join(folder, "it.docx")
            english = os.path.join(folder, "en.docx")
            _narration_docx(italian, [
                ("Prima scena sul Tevere all'alba, con le capanne lontane.", ""),
                ("Seconda scena dentro la capanna, con il fuoco basso.", ""),
            ])
            _narration_docx(english, [
                ("One English scene is enough to keep the other language on disk.", ""),
            ])
            project = create_episode(
                folder, "Conteggi diversi",
                italian_docx=italian,
                english_docx=english,
            )
            self.assertEqual([scene["id"] for scene in load_scenes(project)], [1, 2])
            english_text = Path(narration_path(project, "English")).read_text(encoding="utf-8")
            self.assertIn("One English scene", english_text)

    def test_legacy_doc_and_heading_only_do_not_create_a_folder(self):
        with tempfile.TemporaryDirectory() as folder:
            legacy = os.path.join(folder, "old.doc")
            Path(legacy).write_bytes(b"not a docx")
            italian = os.path.join(folder, "it.docx")
            _narration_docx(italian, [
                ("Il Tevere scorre piano davanti alle capanne all'alba.", ""),
            ])
            with self.assertRaises(ValueError) as raised:
                create_episode(
                    folder, "Da non creare",
                    italian_docx=italian,
                    english_docx=legacy,
                )
            self.assertIn(".docx", str(raised.exception))
            self.assertFalse(os.path.exists(os.path.join(folder, "Da non creare")))

            headings = os.path.join(folder, "headings.docx")
            _narration_docx(headings, [
                ("Solo un titolo", "Title"),
                ("Capitolo", "Heading 1"),
            ])
            with self.assertRaises(ValueError):
                create_episode(folder, "Solo titoli", italian_docx=headings)
            self.assertFalse(os.path.exists(os.path.join(folder, "Solo titoli")))


class VoiceChunkTests(unittest.TestCase):
    def test_sentence_and_paragraph_pauses(self):
        pieces = chunk_narration("Hello. There.\n\nNext paragraph.")
        self.assertEqual(pieces, [
            ("Hello.", 0.8),
            ("There.", 2.5),
            ("Next paragraph.", 0.0),
        ])

    def test_long_sentence_splits_without_a_pause_inside_it(self):
        sentence = "word " * 80
        pieces = chunk_narration(sentence.strip() + ".")
        self.assertGreater(len(pieces), 1)
        self.assertTrue(all(len(text) <= 300 for text, _pause in pieces))
        self.assertTrue(all(pause == 0.0 for _text, pause in pieces[:-1]))

    def test_between_scenes_adds_a_paragraph_pause(self):
        chunks = chunks_for_scene("Hello. There.", trailing_paragraph_pause=True)
        self.assertEqual(chunks[0]["pause_after_s"], 0.8)
        self.assertEqual(chunks[-1]["pause_after_s"], 2.5)

    def test_short_abbreviation_is_not_a_sentence(self):
        pieces = chunk_narration("Rome in c. 509 BC was small.")
        self.assertEqual(len(pieces), 1)
        self.assertIn("c. 509", pieces[0][0])

    def test_missing_reference_fails_before_the_model(self):
        with self.assertRaises(RuntimeError) as raised:
            validate_request("Italian", r"F:\no\such\clip.wav", "Ciao, questa è la mia voce.")
        self.assertIn("reference", str(raised.exception).lower())

    def test_empty_transcript_fails_before_the_model(self):
        with self.assertRaises(RuntimeError) as raised:
            validate_request("English", r"F:\no\such\clip.wav", "  ")
        self.assertIn("transcript", str(raised.exception).lower())

    def test_worker_stays_out_of_the_app_process(self):
        sys.modules.pop("qwen_tts", None)
        import narration.qwen_voice_worker  # noqa: F401
        self.assertNotIn("qwen_tts", sys.modules)
        tree = ast.parse(Path(worker_script()).read_text(encoding="utf-8"))
        imported = []
        for node in tree.body:
            if isinstance(node, ast.Import):
                imported.extend(alias.name for alias in node.names)
            if isinstance(node, ast.ImportFrom):
                imported.append(node.module or "")
        self.assertFalse(any("qwen" in name for name in imported))
        source = Path(worker_script()).read_text(encoding="utf-8")
        self.assertIn("non_streaming_mode=True", source)
        self.assertNotIn("qianwen-res.oss", source)
        app_tree = ast.parse(Path(qwen_voice.__file__).read_text(encoding="utf-8"))
        app_imports = []
        for node in ast.walk(app_tree):
            if isinstance(node, ast.Import):
                app_imports.extend(alias.name for alias in node.names)
            elif isinstance(node, ast.ImportFrom):
                app_imports.append(node.module or "")
        self.assertFalse(any(name.startswith("qwen") for name in app_imports))

    def test_voice_environment_is_the_sibling_venv(self):
        parts = voice_python().parts
        self.assertEqual(parts[-4:], ("voice", "venv", "Scripts", "python.exe"))
        self.assertTrue(voice_python().is_file())
        self.assertTrue(default_model_dir().is_dir())


class SleepyPanelTests(unittest.TestCase):
    def test_panel_has_only_the_chapter_tabs(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PyQt6.QtGui import QColor, QImage
        from PyQt6.QtWidgets import QApplication, QPushButton
        from ui.pipeline_controller import PipelineController
        from ui.sleepy_panel import SleepyPanel, _EpisodeIntakeDialog

        app = QApplication.instance() or QApplication([])
        panel = SleepyPanel(PipelineController())
        dialog = _EpisodeIntakeDialog(panel, profiles={
            "rome_softly_ch01": "Chapter one lock.",
            "rome_softly_ch02": "Chapter two lock.",
        })
        self.assertEqual(dialog.windowTitle(), "Create Sleepy episode")
        self.assertEqual(dialog.values()[:3], ("", "", ""))
        self.assertEqual(dialog.values()[3], "rome_softly_ch01")
        from PyQt6.QtWidgets import QLabel
        dialog_labels = [widget.text() for widget in dialog.findChildren(QLabel)]
        self.assertNotIn("English narration", dialog_labels)
        self.assertIn("Chapter lock for this episode", dialog_labels)
        dialog._profile.setCurrentIndex(1)
        self.assertEqual(dialog.values()[3], "rome_softly_ch02")
        self.assertIn("Chapter two lock.", dialog._lock.toPlainText())
        dialog.deleteLater()
        page_labels = [widget.text() for widget in panel.findChildren(QLabel)]
        self.assertIn("Chapter lock for this episode", page_labels)
        self.assertNotIn("Chapter profile", page_labels)
        titles = [panel.inner_tabs.tabText(index) for index in range(panel.inner_tabs.count())]
        self.assertEqual(
            titles,
            ["Episode", "Script", "Stills", "Lightbox", "Voice", "Final"],
        )
        window_source = Path("ui/main_window.py").read_text(encoding="utf-8")
        self.assertIn('self.mode_tabs.addTab(self.tabs, "Main")', window_source)
        self.assertIn('self.mode_tabs.addTab(self.sleepy_panel, "Sleepy")', window_source)
        panel_source = Path("ui/sleepy_panel.py").read_text(encoding="utf-8")
        self.assertIn("italian_docx", panel_source)
        self.assertIn("profile_key", panel_source)
        self.assertNotIn("english_docx", panel_source)
        self.assertNotIn("QInputDialog", panel_source)
        labels = [button.text() for button in panel.findChildren(QPushButton)]
        self.assertIn("Add English Word", labels)
        self.assertIn("Update Italian Word", labels)
        self.assertIn("Generate this still", labels)
        self.assertIn("Refresh", labels)
        preview = panel._still_preview
        preview.resize(320, 180)
        with tempfile.TemporaryDirectory() as image_dir:
            missing = os.path.join(image_dir, "scene_001_hidream_b01_v1.png")
            preview.show_path(missing)
            self.assertIn("Not painted yet", preview.text())
            painted = os.path.join(image_dir, "scene_002_hidream_b01_v1.png")
            image = QImage(8, 8, QImage.Format.Format_RGB32)
            image.fill(QColor(180, 140, 90))
            self.assertTrue(image.save(painted, "PNG"))
            preview.show_path(painted)
            self.assertFalse(preview.pixmap().isNull())
            self.assertEqual(preview.text(), "")
        panel.deleteLater()
        app.processEvents()


if __name__ == "__main__":
    unittest.main()
